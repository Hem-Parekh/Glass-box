"""Debt vs equity engine: reads a CMA workbook + project inputs, sizes debt from cash flow (DSCR),
and produces numbers for a one-page report. All money in Rs crore unless noted."""
import re, math
import openpyxl

# ---------- 1. CMA parsing ----------
LABELS = {
    "net_sales": [r"^NET SALES"],
    "pbit": [r"^PROFIT BEFORE INTEREST & TAX"],
    "dep": [r"Depreciation"],
    "interest": [r"^Interest & Other Financial Charges"],
    "pbt": [r"^PROFIT BEFORE TAX"],
    "tax": [r"^Current Tax"],
    "pat": [r"^NET PROFIT"],
    "dividend": [r"Dividend Paid \(i\) Amt"],
    "cop": [r"^COST OF PRODUCTION:"],
    "rm": [r"^B\. Indigenous", r"^A\. Imported"],
    "stock_trade": [r"^A\. Stock in Trade"],
    "packing": [r"^B\. Packing Material"],
    "power": [r"Power & Fuel"],
    "labour": [r"Direct Labour"],
    "cost_services": [r"Cost of services"],
    "other_exp": [r"^\(ix\) Other Expenses"],
    "sga": [r"^Selling, General"],
    "net_worth": [r"^NET WORTH"],
    "term_loans": [r"^Term Loans"],
    "tl_current": [r"^Instalments of Term loans"],
    "st_bank": [r"^\(i\) from applicant bank", r"^\(ii\) from other banks"],
    "cash": [r"^Cash and Bank Balances"],
    "liquid1": [r"Govt\. & other securities"],
    "liquid2": [r"Fixed deposits with banks"],
    "tcl": [r"^TOTAL CURRENT LIABILITIES"],
    "ttl": [r"^TOTAL TERM LIABILITIES"],
}

REQUIRED = {"net_sales": "NET SALES", "pbit": "PROFIT BEFORE INTEREST & TAX (PBIT)",
            "pbt": "PROFIT BEFORE TAX", "pat": "NET PROFIT/LOSS (PAT)"}
STRESS_FLOOR = 1.20   # minimum DSCR under stress (EBITDA -20%, rate +2 points)

def _num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

def parse_cma(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    # find header row with years and status row
    years, cols, status = [], [], {}
    for r in range(1, 12):
        vals = [ws.cell(r, c).value for c in range(2, 6)]
        if all(isinstance(v, (int, float)) and 1990 < v < 2100 for v in vals):
            years = [int(v) for v in vals]; cols = list(range(2, 6))
            for c in cols: status[c] = str(ws.cell(r + 1, c).value or "")
            break
    if not years:
        raise ValueError("Could not find the year header row in the CMA sheet")
    data = {}
    for key, pats in LABELS.items():
        tot = [0.0] * len(cols); found = False
        for r in range(1, ws.max_row + 1):
            lab = str(ws.cell(r, 1).value or "").strip()
            if any(re.search(p, lab, re.I) for p in pats):
                found = True
                for i, c in enumerate(cols):
                    tot[i] += _num(ws.cell(r, c).value) or 0.0
                if key not in ("rm", "st_bank"):
                    break
        data[key] = tot if found else [0.0] * len(cols)
        if not found and key in REQUIRED:
            raise ValueError(f"The CMA has no '{REQUIRED[key]}' row in column A of sheet '{ws.title}'")
    if not any(v > 0 for v in data["net_sales"]):
        raise ValueError("The 'NET SALES' row has no positive value in any year column")
    # base year = last column with real sales
    base_i = max(i for i, v in enumerate(data["net_sales"]) if v and v > 0)
    hist = {k: v[:base_i + 1] for k, v in data.items()}
    # growth: CAGR of net sales over available years
    s = hist["net_sales"]
    g = (s[-1] / s[0]) ** (1 / (len(s) - 1)) - 1 if len(s) > 1 and s[0] > 0 else 0.05
    return {"years": years[:base_i + 1], "status": [status[c] for c in cols][:base_i + 1],
            "hist": hist, "base": {k: v[base_i] for k, v in data.items()}, "growth": g}

# ---------- 2. Industry presets (editable defaults) ----------
PRESETS = {
    "manufacturing": dict(min_equity=0.25, dscr_target=1.50, tenor=7, moratorium=1, asset_life=15,
        unit_label="units/yr", note="Asset-heavy: plant & machinery is bankable security, so lenders fund up to ~75% of project cost."),
    "service": dict(min_equity=0.40, dscr_target=1.60, tenor=5, moratorium=0, asset_life=8,
        unit_label="billable units/yr", note="Asset-light: little hard collateral, so lenders expect more promoter equity and stronger cover."),
    "other": dict(min_equity=0.30, dscr_target=1.50, tenor=5, moratorium=0, asset_life=10,
        unit_label="units/yr", note="Trading/retail: funding is mostly working capital; term debt sized on cash flow, WC limits on drawing power."),
}

# ---------- 3. Model ----------
def _schedule(L, rate, tenor, morat):
    n = morat + tenor; bal = L; rows = []
    for t in range(1, n + 1):
        intr = bal * rate
        prin = L / tenor if t > morat else 0.0
        rows.append(dict(year=t, opening=bal, interest=intr, principal=prin, closing=bal - prin))
        bal -= prin
    return rows

def run(cma, p):
    b = cma["base"]; pre = PRESETS[p["industry"]]
    q = {**pre, **{k: v for k, v in p.items() if v is not None}}
    rate, tenor, morat = q["rate"], int(q["tenor"]), int(q["moratorium"])
    n = morat + tenor
    tax_rate = min(max(b["tax"] / b["pbt"], 0.15), 0.35) if b["pbt"] > 0 else 0.25
    g = q.get("growth") if q.get("growth") is not None else cma["growth"]

    # current unit economics (from CMA + current volume)
    units0 = q["existing_capacity"] * q["current_util"]
    if units0 <= 0: raise ValueError("Existing capacity and utilisation must be above zero")
    cop = b["cop"] if b["cop"] > 0 else max(b["net_sales"] - b["pbit"] - b["sga"], 0)
    if q["industry"] == "service":
        var0 = b["cost_services"] + b["labour"]          # people cost scales with billable volume
    else:
        var0 = b["rm"] + b["stock_trade"] + b["packing"] + b["power"]
    if var0 <= 0: var0 = 0.6 * cop                      # fallback when the CMA has no cost split
    b = {**b, "cop": cop}
    price = b["net_sales"] / units0 if units0 else 0
    vc0 = var0 / units0; cpu0 = b["cop"] / units0
    vc_new = vc0 * (1 - q["efficiency_gain"])
    capex = q["project_cost"] - q["wc_margin"]
    sga_ratio = b["sga"] / b["net_sales"] if b["net_sales"] else 0
    dep_p = capex / q["asset_life"]

    def project(L, ebitda_shock=0.0, rate_shock=0.0):
        sch = _schedule(L, rate + rate_shock, tenor, morat)
        out = []
        for i, row in enumerate(sch):
            util = q["ramp"][min(i, len(q["ramp"]) - 1)]
            units = q["new_capacity"] * util
            rev = units * price
            sga_var = rev * sga_ratio * q.get("sga_variable_share", 0.7)   # selling & distribution that scales with sales
            ebitda = (units * (price - vc_new) - sga_var - q["plant_fixed_cost"]) * (1 - ebitda_shock)
            pbt = ebitda - dep_p - row["interest"]
            tax = max(pbt, 0) * tax_rate
            pat = pbt - tax
            cfads = pat + dep_p + row["interest"]
            ds = row["interest"] + row["principal"]
            out.append(dict(**row, util=util, units=units, revenue=rev, ebitda=ebitda, pat=pat,
                            cfads=cfads, ds=ds, dscr=(cfads / ds if ds > 0 else None),
                            fcf_after_ds=cfads - ds))
        return out

    def min_dscr(rows, only_repay=True):
        vals = [r["dscr"] for r in rows if r["dscr"] is not None and (r["principal"] > 0 or not only_repay)]
        return min(vals) if vals else float("inf")

    # company-level cash flow (bank view): existing CFADS grows with sales
    ebitda0 = b["pbit"] + b["dep"]
    def company(L):
        rows = project(L); res = []
        for i, r in enumerate(rows):
            f = (1 + g) ** (i + 1)
            base_cfads = (b["pat"] + b["dep"] + b["interest"]) * f
            existing_ds = b["interest"] + b["tl_current"]
            cf = base_cfads + r["cfads"]
            ds = existing_ds + r["ds"]
            div = b["dividend"] * f
            res.append(dict(year=r["year"], cfads=cf, ds=ds, dscr=cf / ds if ds else None,
                            after_div=cf - div - ds))
        return res

    # solve max project-level debt with DSCR >= target (ring-fenced: the plant must carry its own debt)
    lo, hi = 0.0, q["project_cost"] * 50
    for _ in range(80):
        mid = (lo + hi) / 2
        if min_dscr(project(mid)) >= q["dscr_target"]: lo = mid
        else: hi = mid
    cap_project = lo
    cap_policy = q["project_cost"] * (1 - q["min_equity"])
    debt = min(cap_project, cap_policy)
    debt = math.floor(debt / 5) * 5  # round down to Rs 5 cr
    stress_cut = 0.0
    while debt > 0 and min_dscr(project(debt, ebitda_shock=0.20, rate_shock=0.02)) < STRESS_FLOOR:
        debt -= 5; stress_cut += 5   # cut until the stress case still covers 1.20x
    equity = q["project_cost"] - debt

    rows = project(debt); comp = company(debt)
    stress = project(debt, ebitda_shock=0.20, rate_shock=0.02)
    liquid = b["cash"] + b["liquid1"] + b["liquid2"]
    buffer = q.get("cash_buffer", 0.0)
    internal = max(liquid - buffer, 0)
    new_equity = max(equity - internal, 0)

    ss = rows[-1]
    units_ss = ss["units"]
    cpu_new = (b["cop"] + units_ss * vc_new + q["plant_fixed_cost"] + dep_p) / (units0 + units_ss)
    after_tax_kd = rate * (1 - tax_rate)
    after_tax_cash = q.get("treasury_yield", 0.07) * (1 - tax_rate)
    ke = q.get("cost_of_equity", 0.12)
    base_cfads = b["pat"] + b["dep"] + b["interest"]
    cf_uplift = ss["cfads"] / (base_cfads * (1 + g) ** len(rows))
    binding = "cash flow (DSCR)" if cap_project < cap_policy else "minimum promoter equity"
    if stress_cut: binding = "stress test"
    return dict(q=q, tax_rate=tax_rate, growth=g, price=price, vc0=vc0, vc_new=vc_new, cpu0=cpu0, cpu_new=cpu_new,
                units0=units0, units_ss=units_ss, capex=capex, dep_p=dep_p,
                cap_project=cap_project, cap_policy=cap_policy, debt=debt, equity=equity,
                debt_pct=debt / q["project_cost"], equity_pct=equity / q["project_cost"],
                internal=internal, new_equity=new_equity, rows=rows, company=comp, stress=stress,
                min_dscr=min_dscr(rows), min_dscr_stress=min_dscr(stress),
                min_dscr_company=min(c["dscr"] for c in comp), ebitda0=ebitda0,
                cap_increase=q["new_capacity"] / q["existing_capacity"],
                cum_fcf=sum(r["fcf_after_ds"] for r in rows), ss_cfads=ss["cfads"],
                payback=q["project_cost"] / (sum(r["cfads"] for r in rows) / len(rows)),
                after_tax_kd=after_tax_kd, after_tax_cash=after_tax_cash, ke=ke, binding=binding, stress_cut=stress_cut, liquid=liquid,
                base_cfads=base_cfads, cf_uplift=cf_uplift, sga_ratio=sga_ratio,
                ebitda_margin_ss=ss["ebitda"] / ss["revenue"] if ss["revenue"] else 0)
